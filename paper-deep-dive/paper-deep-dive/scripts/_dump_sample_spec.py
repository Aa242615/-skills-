#!/usr/bin/env python3
"""从 group-meeting-sample.pptx 提取可执行的版式规格：每页形状的坐标、字号、字体、颜色。

用法: python scripts/_dump_sample_spec.py
输出: references/_sample_spec.txt
"""
import sys, os
from pptx import Presentation
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
PPTX = os.path.join(ROOT, "references", "group-meeting-sample.pptx")
OUT = os.path.join(ROOT, "references", "_sample_spec.txt")

def emu_in(v):
    try:
        return round(Emu(int(v)).inches, 3)
    except Exception:
        return None

def font_of(run):
    rPr = run._r.find("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")
    latin = run.font.name
    ea = cs = None
    if rPr is not None:
        ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        for tag in ("a:ea", "a:cs"):
            el = rPr.find(f"{ns}{tag}")
            if el is not None:
                v = el.get("typeface")
                if tag == "a:ea": ea = v
                else: cs = v
    return latin, ea, cs

def main():
    prs = Presentation(PPTX)
    lines = [f"# 样例版式规格 dump：{os.path.basename(PPTX)}",
             f"画布 {emu_in(prs.slide_width)} x {emu_in(prs.slide_height)} 英寸，共 {len(prs.slides)} 页", ""]
    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"===== 第 {i} 页 =====")
        for sh in slide.shapes:
            kind = str(sh.shape_type) if sh.shape_type is not None else "?"
            info = [f"[{kind}] name={sh.name!r}"]
            for attr in ("left", "top", "width", "height"):
                v = getattr(sh, attr, None)
                info.append(f"{attr}={emu_in(v)}")
            info.append(f"rot={getattr(sh, 'rotation', None)}")
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    runs = []
                    for r in p.runs:
                        sz = r.font.size.pt if r.font.size else None
                        col = None
                        try:
                            if r.font.color and r.font.color.type is not None and r.font.color.rgb is not None:
                                col = str(r.font.color.rgb)
                        except Exception:
                            pass
                        lat, ea, cs = font_of(r)
                        runs.append(f"({r.text[:40]!r} {sz}pt b={r.font.bold} latin={lat} ea={ea} cs={cs} color={col})")
                    if runs:
                        info.append("P[" + " | ".join(runs) + "]")
            if sh.has_table:
                info.append(f"TABLE {len(sh.table.rows)}x{len(sh.table.columns)}")
            lines.append("  " + " ".join(info))
        lines.append("")
    with open(OUT, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"已写入 {OUT}（{len(lines)} 行）")

if __name__ == "__main__":
    main()
